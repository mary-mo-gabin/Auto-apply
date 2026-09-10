import pytest
from pathlib import Path
from docx import Document

import document_utils

def test_get_master_resume_text_success(tmp_path, monkeypatch):
    """Tests if the function correctly reads and extracts text from a .docx file."""
    # Arrange: Create a temporary docx file with fake resume data
    fake_resume_path = tmp_path / "master_resume.docx"
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Full-Stack Software Engineer")
    doc.save(fake_resume_path)
    
    # Hijack the global variable in the module to point to the fake file
    monkeypatch.setattr("document_utils.MASTER_RESUME_PATH", fake_resume_path)
    
    # Act: Run the function
    result = document_utils.get_master_resume_text()
    
    # Assert: Verify the extracted text matches
    assert "Jane Doe" in result
    assert "Full-Stack Software Engineer" in result
    
def test_get_master_resume_text_missing_file(tmp_path, monkeypatch):
    """Tests if the function correctly raises an error when the file is missing."""
    # Arrange: Point the global variable to a file that does not exist
    fake_missing_path = tmp_path / "does_not_exist.docx"
    monkeypatch.setattr("document_utils.MASTER_RESUME_PATH", fake_missing_path)
    
    # Act & Assert: Prove that it raises a FileNotFoundError
    with pytest.raises(FileNotFoundError, match="Master resume not found"):
        document_utils.get_master_resume_text()